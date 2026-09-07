from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "Bao_cao_do_an_ComZone.docx"
OUTPUT = ROOT / "deliverables" / "Bao_cao_do_an_ComZone_co_hinh_anh.docx"
SHOT_DIR = ROOT / "deliverables" / "app_screenshots"


FIGURES = [
    (
        "5.9.1. Đăng nhập hệ thống",
        "01_dang_nhap.png",
        "Màn hình đăng nhập ComZone",
        "Chức năng đăng nhập tiếp nhận email và mật khẩu, hỗ trợ ghi nhớ trạng thái phiên và điều hướng sang đăng ký hoặc khôi phục mật khẩu. Khi xác thực thành công, hệ thống đọc vai trò của tài khoản để hiển thị đúng khu vực chức năng.",
        "Biểu mẫu đăng nhập gồm trường email, mật khẩu và các liên kết hỗ trợ xác thực.",
    ),
    (
        "5.9.2. Đăng ký tài khoản",
        "02_dang_ky.png",
        "Màn hình đăng ký tài khoản",
        "Người dùng mới có thể tạo tài khoản bằng họ tên, email, số điện thoại và mật khẩu. Dữ liệu đầu vào được kiểm tra trước khi gửi đến Firebase Authentication và tạo hồ sơ người dùng trong Firestore.",
        "Biểu mẫu đăng ký tài khoản người dùng mới trên ComZone.",
    ),
    (
        "5.9.3. Trang chủ và danh mục truyện",
        "03_trang_chu_danh_muc.png",
        "Trang chủ và khu vực khám phá truyện",
        "Trang chủ giới thiệu nền tảng, cung cấp ô tìm kiếm, lối tắt đến truyện tranh và đấu giá, đồng thời hiển thị danh mục sản phẩm. Thanh điều hướng dưới thay đổi theo vai trò; tài khoản quản trị trong lần chạy thử có thể truy cập thêm Seller Hub và trang quản trị.",
        "Trang chủ sau khi đăng nhập bằng tài khoản quản trị, có đầy đủ thanh điều hướng theo vai trò.",
    ),
    (
        "5.9.4. Xem chi tiết truyện",
        "04_chi_tiet_truyen.png",
        "Màn hình chi tiết truyện",
        "Màn hình chi tiết trình bày ảnh bìa, tên truyện, nhà xuất bản, thể loại, tình trạng, số lượng còn lại, giá bán và đánh giá. Người dùng có thể xem hồ sơ người bán, lưu truyện hoặc thêm sản phẩm vào giỏ hàng.",
        "Chi tiết một truyện đang được đăng bán trên hệ thống.",
    ),
    (
        "5.9.5. Danh sách phiên đấu giá",
        "05_dau_gia.png",
        "Màn hình danh sách phiên đấu giá",
        "Khu vực đấu giá phân loại phiên theo ba trạng thái: đang diễn ra, sắp diễn ra và đã kết thúc. Tại thời điểm chụp, bộ lọc đang diễn ra chưa có phiên phù hợp nên ứng dụng hiển thị trạng thái rỗng thay vì báo lỗi.",
        "Danh sách phiên đấu giá với bộ lọc trạng thái và thông báo khi chưa có dữ liệu.",
    ),
    (
        "5.9.6. Trao đổi truyện trong cộng đồng",
        "06_trao_doi.png",
        "Màn hình trao đổi truyện",
        "Chức năng trao đổi tập hợp các bài đăng cộng đồng và cung cấp nút tạo bài cho người dùng đã đăng nhập. Ảnh chụp thể hiện giao diện trạng thái rỗng khi chưa có bài trao đổi đang hoạt động.",
        "Khu vực trao đổi cộng đồng và thao tác tạo bài đăng mới.",
    ),
    (
        "5.9.7. Seller Hub",
        "07_seller_hub.png",
        "Bảng điều khiển Seller Hub",
        "Seller Hub tổ chức nghiệp vụ người bán theo các thẻ Dashboard, Đơn hàng, Kho truyện, Đấu giá và Payout. Tài khoản thử nghiệm chưa phát sinh dữ liệu Seller tại thời điểm chụp nên hệ thống hiển thị thông báo phù hợp.",
        "Seller Hub với các nhóm chức năng quản lý hoạt động bán hàng.",
    ),
    (
        "5.9.8. Tổng quan quản trị",
        "08_quan_tri.png",
        "Trang tổng quan quản trị",
        "Trang quản trị tổng hợp số đơn đặt mua, tin đăng và thành viên theo dữ liệu Firestore. Thanh menu bên trái cung cấp các mô-đun quản lý đơn hàng, tin đăng, đấu giá, phản hồi, ví, duyệt Seller, người dùng, danh mục và cài đặt.",
        "Bảng điều khiển tổng quan dành cho tài khoản có vai trò quản trị.",
    ),
    (
        "5.9.9. Quản lý đơn hàng",
        "09_quan_ly_don_hang.png",
        "Màn hình quản lý đơn hàng",
        "Quản trị viên có thể tìm kiếm, lọc và theo dõi trạng thái đơn hàng. Giao diện mô tả quy trình Đơn mới – Phân loại – Gom nhóm – Xử lý, đồng thời hiển thị thông tin truyện, người mua, phương thức thanh toán, mức ưu tiên và trạng thái xử lý.",
        "Danh sách và quy trình xử lý đơn hàng trong trang quản trị.",
    ),
    (
        "5.9.10. Theo dõi đơn hàng cá nhân",
        "13_don_hang_ca_nhan.png",
        "Màn hình đơn hàng của người dùng",
        "Người mua theo dõi các đơn đã đặt theo tên truyện, thời điểm tạo, giá trị và trạng thái hiện tại. Mỗi dòng có thể được mở để xem thông tin chi tiết và tiến trình xử lý.",
        "Danh sách đơn hàng thuộc tài khoản đang đăng nhập.",
    ),
    (
        "5.9.11. Đối soát ví",
        "10_quan_ly_vi.png",
        "Màn hình quản lý yêu cầu ví",
        "Mô-đun ví hỗ trợ quản trị viên đối soát yêu cầu nạp và rút tiền, bao gồm số tiền, cổng thanh toán, mã người dùng, thời điểm và trạng thái giao dịch. Ảnh chụp cho thấy một giao dịch nạp tiền qua VNPay đã hoàn tất.",
        "Danh sách yêu cầu ví và trạng thái đối soát trong trang quản trị.",
    ),
    (
        "5.9.12. Thông báo",
        "11_thong_bao.png",
        "Màn hình thông báo",
        "Màn hình thông báo tập trung các cập nhật liên quan đến tài khoản, đơn hàng, đấu giá và trao đổi. Khi chưa có thông báo, hệ thống hiển thị trạng thái rỗng rõ ràng để người dùng không nhầm với lỗi tải dữ liệu.",
        "Trạng thái màn hình thông báo khi tài khoản chưa có dữ liệu.",
    ),
    (
        "5.9.13. Hồ sơ cá nhân",
        "12_ho_so_ca_nhan.png",
        "Màn hình hồ sơ cá nhân",
        "Người dùng có thể xem và cập nhật họ tên, số điện thoại; email được khóa vì là định danh đăng nhập. Nút Lưu hồ sơ ghi lại các thay đổi hợp lệ vào hồ sơ người dùng trên Firestore.",
        "Biểu mẫu xem và cập nhật thông tin hồ sơ cá nhân.",
    ),
]


def set_run_font(run, size: float = 13, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if italic is not None:
        run.italic = italic


def style_body(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        set_run_font(run)


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def move_before(paragraph, target_paragraph) -> None:
    target_paragraph._p.addprevious(paragraph._p)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    for _, filename, *_ in FIGURES:
        path = SHOT_DIR / filename
        if not path.exists():
            raise FileNotFoundError(path)

    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)

    target = next(
        p for p in document.paragraphs if p.text.strip() == "CHƯƠNG 6. KIỂM THỬ VÀ ĐÁNH GIÁ"
    )
    created = []

    section_heading = document.add_paragraph(
        "5.9. Hình ảnh minh họa các chức năng", style="Heading 2"
    )
    section_heading.paragraph_format.page_break_before = True
    section_heading.paragraph_format.keep_with_next = True
    created.append(section_heading)

    intro = document.add_paragraph(
        "Các hình dưới đây được chụp trực tiếp từ phiên bản Flutter Web của ComZone khi chạy thử trên máy phát triển ngày 04/09/2026. Dữ liệu hiển thị là dữ liệu thử nghiệm thực tế tại thời điểm chụp; các màn hình không phát sinh dữ liệu được giữ nguyên trạng thái rỗng để phản ánh đúng hành vi của hệ thống."
    )
    style_body(intro)
    created.append(intro)

    first = True
    figure_number = 6
    for heading, filename, caption, description, alt_description in FIGURES:
        h = document.add_paragraph(heading, style="Heading 3")
        h.paragraph_format.keep_with_next = True
        if not first:
            h.paragraph_format.page_break_before = True
        created.append(h)

        body = document.add_paragraph(description)
        style_body(body)
        body.paragraph_format.keep_with_next = True
        created.append(body)

        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.keep_with_next = True
        image_run = image_paragraph.add_run()
        inline_shape = image_run.add_picture(str(SHOT_DIR / filename), width=Cm(15.6))
        set_alt_text(inline_shape, caption, alt_description)
        created.append(image_paragraph)

        caption_paragraph = document.add_paragraph(style="Caption")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_paragraph.add_run(f"Hình {figure_number}. {caption}")
        set_run_font(caption_run, size=12, italic=True)
        caption_paragraph.paragraph_format.space_after = Pt(6)
        created.append(caption_paragraph)

        figure_number += 1
        first = False

    for paragraph in created:
        move_before(paragraph, target)

    document.core_properties.title = "Báo cáo đồ án ComZone có hình ảnh minh họa chức năng"
    document.core_properties.subject = "Báo cáo đồ án cơ sở - ứng dụng mua bán, đấu giá và trao đổi truyện tranh"
    document.core_properties.keywords = "ComZone, Flutter, Firebase, đồ án, ảnh chụp chức năng"

    # Yêu cầu Word cập nhật TOC/field ngay khi người dùng mở tài liệu.
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
